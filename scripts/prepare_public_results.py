#!/usr/bin/env python3
"""Extract aggregate-only public result tables from final manuscript artifacts.

This is a one-time release-building utility. It reads the author-designated
final DOCX files and audited aggregate outputs, never writes to the source
project, and emits only tables or histograms that do not contain row-level
coordinates, timestamps, targets, or fused records.
"""

from __future__ import annotations

import argparse
import csv
import re
import shutil
from pathlib import Path

import numpy as np
from docx import Document


def clean(value: str) -> str:
    return " ".join(value.split())


def table_rows(table) -> list[list[str]]:
    return [[clean(cell.text) for cell in row.cells] for row in table.rows]


def write_rows(path: Path, rows: list[list[str]], extra_header: str | None = None, extra_value: str | None = None) -> None:
    if not rows:
        raise ValueError(f"No rows for {path}")
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    if extra_header is not None:
        normalized[0] = [extra_header] + normalized[0]
        for index in range(1, len(normalized)):
            normalized[index] = [extra_value or ""] + normalized[index]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        csv.writer(stream).writerows(normalized)


def combine_tables(path: Path, sections: list[tuple[str, object]]) -> None:
    all_rows: list[list[str]] = []
    canonical_header: list[str] | None = None
    for section, table in sections:
        rows = table_rows(table)
        header, body = rows[0], rows[1:]
        if canonical_header is None:
            canonical_header = header
            all_rows.append(["section"] + header)
        elif header != canonical_header:
            raise ValueError(f"Header mismatch while creating {path}: {header} != {canonical_header}")
        all_rows.extend([[section] + row for row in body])
    write_rows(path, all_rows)


def extract_main_tables(manuscript: Document, staging: Path) -> None:
    results = staging / "results"
    write_rows(results / "main_table_1_dataset_roles.csv", table_rows(manuscript.tables[0]))

    meanstate = table_rows(manuscript.tables[28])
    meanstate_tidy = [["evaluation_split"] + meanstate[1]]
    meanstate_tidy.extend([["2022_model_selection_validation"] + row for row in meanstate[2:8]])
    meanstate_tidy.extend([["2023_temporal_holdout"] + row for row in meanstate[10:16]])
    write_rows(results / "main_table_2_meanstate_performance.csv", meanstate_tidy)

    anomaly = table_rows(manuscript.tables[29])
    anomaly_tidy = [["evaluation_split"] + anomaly[1]]
    anomaly_tidy.extend([["2022_model_selection_validation"] + row for row in anomaly[2:5]])
    anomaly_tidy.extend([["2023_temporal_holdout"] + row for row in anomaly[7:10]])
    write_rows(results / "main_table_3_daily_anomaly_performance.csv", anomaly_tidy)

    write_rows(results / "main_table_4_bounded_adjustment.csv", table_rows(manuscript.tables[30]))


def extract_si_tables(si: Document, staging: Path) -> None:
    results = staging / "results"
    write_rows(results / "si_table_s1_dataset_partitions.csv", table_rows(si.tables[0]))
    write_rows(results / "si_table_s2_inventory_support_overlap.csv", table_rows(si.tables[1]))
    combine_tables(
        results / "si_table_s3_monthly_meanstate_r2.csv",
        [("2022_validation", si.tables[2]), ("2023_external", si.tables[3])],
    )
    combine_tables(
        results / "si_table_s4_seasonal_regional_meanstate_r2.csv",
        [("season", si.tables[4]), ("region", si.tables[5])],
    )
    combine_tables(
        results / "si_table_s5_grouped_daily_anomaly_r2.csv",
        [
            ("2022_month", si.tables[6]),
            ("2022_region", si.tables[7]),
            ("2023_month", si.tables[8]),
            ("2023_season", si.tables[9]),
            ("2023_region", si.tables[10]),
        ],
    )
    write_rows(results / "si_table_s6_bounded_adjustment_by_month.csv", table_rows(si.tables[11]))
    write_rows(results / "si_table_s7a_area_weighted_performance.csv", table_rows(si.tables[12]))
    write_rows(results / "si_table_s7b_area_weighted_contrasts.csv", table_rows(si.tables[13]))
    write_rows(results / "si_table_s8_geographic_holdout.csv", table_rows(si.tables[14]))
    combine_tables(
        results / "si_table_s9_monthly_sampling_support.csv",
        [("2022", si.tables[15]), ("2023", si.tables[16])],
    )
    write_rows(results / "si_table_s10_common_mask_audit.csv", table_rows(si.tables[17]))
    write_rows(results / "si_table_s11_burden_scale_errors.csv", table_rows(si.tables[18]))
    write_rows(results / "si_table_s12_bccmass_provenance.csv", table_rows(si.tables[19]))


def copy_for_figure_sources(staging: Path) -> None:
    pairs = {
        "results/main_table_2_meanstate_performance.csv": "figure_source_data/figure_2_meanstate_performance.csv",
        "results/main_table_3_daily_anomaly_performance.csv": "figure_source_data/figure_4_daily_anomaly_performance.csv",
        "results/main_table_4_bounded_adjustment.csv": "figure_source_data/figure_5_bounded_adjustment.csv",
        "results/si_table_s9_monthly_sampling_support.csv": "figure_source_data/figure_s1_monthly_sampling_support.csv",
    }
    for source_name, destination_name in pairs.items():
        destination = staging / destination_name
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(staging / source_name, destination)


def make_figure1_workflow(staging: Path) -> None:
    rows = [
        ["branch", "inputs", "target_or_output", "reported_elements"],
        ["mean_state", "inventory; monthly meteorology; controls", "monthly mean log burden", "Table 2; Figures 2-3"],
        ["daily_anomaly", "meteorological anomalies; fire or AEF context", "daily log-burden anomaly", "Table 3; Figure 4"],
        ["bounded_adjustment", "frozen joint mean-state model; channel offsets", "reconstruction-deviation trade-off", "Table 4; Figure 5"],
    ]
    write_rows(staging / "figure_source_data" / "figure_1_workflow.csv", rows)


def make_figure3_summary(manuscript: Document, staging: Path) -> None:
    text = " ".join(clean(paragraph.text) for paragraph in manuscript.paragraphs)
    match = re.search(
        r"lower local absolute error in\s+([0-9.]+)%\s+of valid cells in November and\s+([0-9.]+)%\s+in December",
        text,
    )
    if not match:
        raise ValueError("Could not locate Figure 3 summary percentages")
    rows = [
        ["month", "joint_lower_absolute_error_fraction_percent", "public_reproduction_scope"],
        ["2022-11", match.group(1), "aggregate summary only; pointwise/cellwise source arrays are available only through the restricted-access record, subject to the approved access conditions."],
        ["2022-12", match.group(2), "aggregate summary only; pointwise/cellwise source arrays are available only through the restricted-access record, subject to the approved access conditions."],
    ]
    write_rows(staging / "figure_source_data" / "figure_3_spatial_summary_only.csv", rows)


def make_sampling_histograms(npz_path: Path, staging: Path) -> None:
    archive = np.load(npz_path, allow_pickle=False)
    rows = [["year_month", "valid_day_count", "grid_cell_count", "grid_cell_fraction"]]
    denominator: int | None = None
    for key in sorted(name for name in archive.files if name.endswith("_i380_finite_count")):
        array = archive[key]
        if array.ndim != 2:
            raise ValueError(f"Unexpected map dimensionality for {key}: {array.shape}")
        denominator = int(array.size)
        counts = np.bincount(array.astype(np.int64).ravel())
        year_month = key.split("_", 1)[0]
        for day_count, cell_count in enumerate(counts.tolist()):
            rows.append([year_month, str(day_count), str(cell_count), f"{cell_count / denominator:.12g}"])
    if denominator != 361 * 576:
        raise ValueError(f"Unexpected grid size: {denominator}")
    write_rows(staging / "figure_source_data" / "figure_s2_valid_day_count_histograms.csv", rows)


def make_sampling_spatial_summary(npz_path: Path, staging: Path) -> None:
    archive = np.load(npz_path, allow_pickle=False)
    rows = [["year", "included_months", "mismatch_cell_months", "maximum_mismatch_months_at_any_cell", "public_reproduction_scope"]]
    for year in (2022, 2023):
        keys = sorted(name for name in archive.files if name.startswith(f"{year}-") and name.endswith("_eval_difference_class"))
        mismatch_sum = None
        total_mismatch = 0
        for key in keys:
            values = archive[key]
            mismatch = np.logical_or(values == 1, values == 2)
            total_mismatch += int(mismatch.sum())
            mismatch_sum = mismatch.astype(np.uint16) if mismatch_sum is None else mismatch_sum + mismatch
        maximum = int(mismatch_sum.max()) if mismatch_sum is not None else 0
        rows.append(
            [
                str(year), str(len(keys)), str(total_mismatch), str(maximum),
                "aggregate mismatch summary only; pointwise/cellwise source arrays are available only through the restricted-access record, subject to the approved access conditions.",
            ]
        )
    write_rows(staging / "figure_source_data" / "figure_s3_spatial_common_support_summary_only.csv", rows)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manuscript-docx", type=Path, required=True)
    parser.add_argument("--si-docx", type=Path, required=True)
    parser.add_argument("--sampling-maps-npz", type=Path, required=True)
    parser.add_argument("--staging-root", type=Path, required=True)
    args = parser.parse_args()

    staging = args.staging_root.resolve()
    manuscript = Document(str(args.manuscript_docx))
    si = Document(str(args.si_docx))
    if len(manuscript.tables) < 31 or len(si.tables) < 20:
        raise ValueError("Unexpected manuscript/SI table count; refusing ambiguous extraction")

    extract_main_tables(manuscript, staging)
    extract_si_tables(si, staging)
    copy_for_figure_sources(staging)
    make_figure1_workflow(staging)
    make_figure3_summary(manuscript, staging)
    make_sampling_histograms(args.sampling_maps_npz, staging)
    make_sampling_spatial_summary(args.sampling_maps_npz, staging)
    print("PUBLIC_RESULT_TABLES=17")
    print("FIGURE_SOURCE_TABLES=8")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
